from datetime import date, timedelta
from pathlib import Path
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from core.models import Task, LearningPlan, Phase, Module, Lesson

PRI = {"Essential": "high", "Important": "medium", "Optional": "low"}

class Command(BaseCommand):
    help = "Import the '60-Day Plan' sheet of the tracker workbook as tasks for one user (idempotent, never overwrites progress)."
    def add_arguments(self, p):
        p.add_argument("username"); p.add_argument("--start", required=True, help="Day 1 date, YYYY-MM-DD")
        p.add_argument("--file", default="data/Faiz_Recovery_Tracker.xlsx"); p.add_argument("--dry-run", action="store_true")
    def handle(self, *a, username, start, file, dry_run, **kw):
        user = get_user_model().objects.filter(username=username).first()
        if not user: raise CommandError(f"User '{username}' does not exist. Register first.")
        if not Path(file).exists(): raise CommandError(f"File not found: {file}")
        d1 = date.fromisoformat(start)
        ws = load_workbook(file, data_only=True)["60-Day Plan"]
        rows, errors = [], []
        for r in range(2, ws.max_row + 1):
            day = ws.cell(r, 1).value
            if not isinstance(day, int): continue
            title = ws.cell(r, 6).value
            if not title: errors.append(f"row {r}: empty topic"); continue
            hrs = ws.cell(r, 11).value or 1
            rows.append(dict(plan_day=day, title=str(title)[:300], subject=str(ws.cell(r, 4).value or "")[:80], phase=str(ws.cell(r, 3).value or "")[:80],
                priority=PRI.get(ws.cell(r, 5).value, "medium"), est_minutes=int(float(hrs) * 60),
                resource=str(ws.cell(r, 12).value or "")[:300], scheduled_date=d1 + timedelta(days=day - 1),
                description=f"Exercise: {ws.cell(r, 7).value}\nDSA: {ws.cell(r, 8).value}\nDone when: {ws.cell(r, 13).value}"))
        for e in errors: self.stderr.write(e)
        if errors: raise CommandError("Fix the errors above; nothing imported.")
        if dry_run: self.stdout.write(f"Dry run OK: {len(rows)} tasks would be considered."); return
        created = 0
        with transaction.atomic():
            for row in rows:
                _, c = Task.objects.get_or_create(user=user, plan_day=row["plan_day"], defaults=row); created += c
            # Keep the workbook tasks as-is while creating an editable curriculum tree once.
            plan, made_plan = LearningPlan.objects.get_or_create(owner=user, title="Imported 60-Day Plan",
                defaults={"description":"Curriculum imported from the Faiz Recovery Tracker workbook.", "start_date":d1,
                          "target_date":d1 + timedelta(days=59)})
            if made_plan:
                phases = {}
                modules = {}
                for row in rows:
                    phase_title = row["phase"] or "General"
                    if phase_title not in phases:
                        phases[phase_title] = Phase.objects.create(plan=plan, title=phase_title, position=len(phases))
                    key = (phase_title, row["subject"] or "General")
                    if key not in modules:
                        modules[key] = Module.objects.create(phase=phases[phase_title], title=key[1], position=len([x for x in modules if x[0] == phase_title]))
                    Lesson.objects.create(module=modules[key], title=row["title"], description=row["description"], resource=row["resource"],
                                          est_minutes=row["est_minutes"], position=Lesson.objects.filter(module=modules[key]).count())
        self.stdout.write(f"Imported {created} new tasks ({len(rows) - created} already existed); {'created editable curriculum' if made_plan else 'curriculum already existed'}.")
